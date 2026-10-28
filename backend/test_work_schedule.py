import asyncio
import itertools
import os
import sys
import types
import unittest
from datetime import date, datetime, timedelta
from unittest import mock
from zoneinfo import ZoneInfo

from work_schedule_rules import (
    DEFAULT_INTERN_WORK_SCHEDULE,
    is_start_within_schedule,
    time_on_day,
    validate_work_schedule,
)


def _import_main():
    """Importa o main.py para testar as regras puras dele.

    O fastapi_mail so serve para mandar e-mail e costuma faltar na maquina de
    desenvolvimento; sem ele, entra um stub minimo. Faltando outra dependencia
    (fastapi, pydantic...), os testes que usam o main sao pulados.
    """
    try:
        import main
        return main
    except ImportError as exc:
        if exc.name != "fastapi_mail":
            return None
    stub = types.ModuleType("fastapi_mail")

    class _Any:
        def __init__(self, *args, **kwargs):
            pass

    stub.FastMail = stub.MessageSchema = stub.ConnectionConfig = _Any
    stub.MessageType = types.SimpleNamespace(plain="plain", html="html")
    sys.modules["fastapi_mail"] = stub
    try:
        import main
        return main
    except ImportError:
        return None


main = _import_main()
needs_main = unittest.skipIf(main is None, "dependencias do main.py ausentes")


def _call(result):
    """Chama o endpoint como o FastAPI: os que so fazem trabalho bloqueante sao
    `def` (rodam fora do event loop); os que esperam HTTP sao `async def`."""
    return asyncio.run(result) if asyncio.iscoroutine(result) else result


class WorkScheduleRulesTest(unittest.TestCase):
    def setUp(self):
        self.schedule = validate_work_schedule(DEFAULT_INTERN_WORK_SCHEDULE)
        self.timezone = ZoneInfo("America/Sao_Paulo")

    def at(self, hour: int, minute: int) -> datetime:
        return datetime(2026, 7, 17, hour, minute, tzinfo=self.timezone)

    def test_start_window_boundaries_are_inclusive(self):
        self.assertTrue(is_start_within_schedule(self.at(13, 25), self.schedule))
        self.assertTrue(is_start_within_schedule(self.at(19, 25), self.schedule))

    def test_start_outside_window_is_blocked(self):
        self.assertFalse(is_start_within_schedule(self.at(13, 24), self.schedule))
        self.assertFalse(is_start_within_schedule(self.at(19, 26), self.schedule))

    def test_weekend_start_always_requires_approval(self):
        saturday = datetime(2026, 7, 18, 13, 30, tzinfo=self.timezone)
        sunday = datetime(2026, 7, 19, 13, 30, tzinfo=self.timezone)
        self.assertFalse(is_start_within_schedule(saturday, self.schedule))
        self.assertFalse(is_start_within_schedule(sunday, self.schedule))

    def test_default_overtime_threshold_is_1935(self):
        threshold = time_on_day(
            self.at(18, 0).date(),
            self.schedule["overtime_prompt_after"],
            self.schedule["timezone"],
        )
        self.assertEqual(threshold, self.at(19, 35))

    def test_invalid_schedule_is_rejected(self):
        with self.assertRaises(Exception):
            validate_work_schedule({"shift_start": "25:00"})

    def test_single_digit_hours_are_zero_padded(self):
        # Como texto "9:30" > "13:30"; sem normalizar, um turno valido da
        # manha era recusado como "inicio depois do fim".
        schedule = validate_work_schedule({
            "shift_start": "9:30",
            "shift_end": "13:30",
            "start_allowed_from": "9:25",
            "start_allowed_until": "13:25",
            "overtime_prompt_after": "13:35",
        })
        self.assertEqual(schedule["shift_start"], "09:30")
        self.assertEqual(schedule["start_allowed_from"], "09:25")
        self.assertEqual(schedule["shift_end"], "13:30")

    def test_padded_comparison_rejects_inverted_shift(self):
        # Como texto "14:00" < "9:00" e o turno invertido passava.
        with self.assertRaisesRegex(ValueError, "inicio do turno"):
            validate_work_schedule({
                "shift_start": "14:00",
                "shift_end": "9:00",
                "overtime_prompt_after": "9:05",
            })


@needs_main
class HolidayKeysTest(unittest.TestCase):
    def data(self, holidays, start=None, end=None, month="10", year="2026"):
        return main.ExportData(
            employee_name="Teste",
            employee_role="intern",
            month=month,
            year=year,
            is_campo=False,
            daily_records=[],
            holidays=holidays,
            period_start=start,
            period_end=end,
        )

    def test_calendar_month_uses_day_of_month(self):
        keys = main.holiday_date_keys(self.data([12]))
        self.assertEqual(keys, {"2026-10-12"})

    def test_custom_period_does_not_mark_position_in_period(self):
        # 29/09..28/10: o 15o dia do periodo e 13/10, que nao e feriado.
        keys = main.holiday_date_keys(
            self.data([12, 15], "2026-09-29T00:00:00", "2026-10-28T23:59:59")
        )
        self.assertEqual(keys, {"2026-10-12", "2026-10-15"})

    def test_custom_period_matches_day_in_first_month(self):
        keys = main.holiday_date_keys(
            self.data([30], "2026-09-29T00:00:00", "2026-10-28T23:59:59")
        )
        self.assertEqual(keys, {"2026-09-30"})

    def test_holiday_keys_reach_intern_payroll(self):
        record = main.ExportRecord(date="2026-10-12", check_in="13:30", check_out="19:30")
        keys = main.holiday_date_keys(self.data([12]))
        payroll = main.classify_estag_payroll(record, date(2026, 10, 12), keys, None)
        self.assertEqual(payroll["100"], 360)

    def test_day_after_period_uses_national_calendar(self):
        # Dezembro: o "1" digitado seria 01/12; 01/01 vem do calendario.
        keys = main.holiday_date_keys(self.data([25], month="12"))
        self.assertEqual(keys, {"2026-12-25", "2027-01-01"})
        record = main.ExportRecord(
            date="2026-12-31", check_in="19:30", check_out="20:00", break_start="21:00", break_end="07:00"
        )
        payroll = main.classify_estag_payroll(record, date(2026, 12, 31), keys, None)
        self.assertEqual(payroll["100"], 540)
        self.assertEqual(payroll["normal"], 90)
        # Abril: 01/05.
        self.assertIn("2026-05-01", main.holiday_date_keys(self.data([], month="04")))
        # 29/02..28/03/2024: o dia seguinte e a Sexta-feira Santa.
        keys = main.holiday_date_keys(
            self.data([], "2024-02-29T00:00:00", "2024-03-28T23:59:59", month="03", year="2024")
        )
        self.assertEqual(keys, {"2024-03-29"})
        # Dia seguinte comum nao entra.
        self.assertEqual(main.holiday_date_keys(self.data([])), set())

    def test_national_holiday_calendar(self):
        self.assertTrue(main.is_national_holiday(date(2026, 4, 3)))  # Pascoa 05/04/2026
        self.assertTrue(main.is_national_holiday(date(2024, 11, 20)))
        self.assertFalse(main.is_national_holiday(date(2023, 11, 20)))
        self.assertFalse(main.is_national_holiday(date(2026, 10, 15)))


@needs_main
class InternSaturdayPayrollTest(unittest.TestCase):
    saturday = date(2026, 10, 17)

    def classify(self, check_in, check_out, **rules):
        record = main.ExportRecord(date=self.saturday.isoformat(), check_in=check_in, check_out=check_out)
        return main.classify_estag_payroll(record, self.saturday, set(), main.EstagRules(**rules))

    def test_saturday_night_goes_to_100_and_rest_to_50(self):
        payroll = self.classify("18:00", "23:30")
        self.assertEqual(payroll["100"], 90)
        self.assertEqual(payroll["50"], 240)
        self.assertEqual(payroll["normal"] + payroll["extras"], 0)

    def test_saturday_without_night_rule_is_all_50(self):
        payroll = self.classify("18:00", "23:30", night_100=False)
        self.assertEqual(payroll["50"], 330)
        self.assertEqual(payroll["100"], 0)

    def test_saturday_daytime_is_all_50(self):
        payroll = self.classify("08:00", "12:00")
        self.assertEqual(payroll["50"], 240)
        self.assertEqual(payroll["100"], 0)

    def test_saturday_without_50_follows_weekday_rule(self):
        payroll = self.classify("18:00", "23:30", saturday_50=False)
        self.assertEqual(payroll["100"], 90)
        self.assertEqual(payroll["normal"], 240)
        self.assertEqual(payroll["50"], 0)


@needs_main
class InternPayrollAcrossMidnightTest(unittest.TestCase):
    def classify(self, day, holidays=(), **columns):
        record = main.ExportRecord(date=day.isoformat(), **columns)
        return main.classify_estag_payroll(record, day, set(holidays), None)

    def test_saturday_night_into_sunday(self):
        # Como o app: domingo 05-07 e 100%, nao 50% pela data da linha.
        payroll = self.classify(
            date(2026, 10, 17), check_in="19:30", check_out="20:00", break_start="21:00", break_end="07:00"
        )
        self.assertEqual(payroll["50"], 90)
        self.assertEqual(payroll["100"], 540)
        self.assertEqual(payroll["normal"] + payroll["extras"], 0)

    def test_friday_night_into_saturday(self):
        payroll = self.classify(
            date(2026, 10, 16), check_in="19:30", check_out="20:00", break_start="21:00", break_end="07:00"
        )
        self.assertEqual(payroll["normal"], 90)
        self.assertEqual(payroll["100"], 420)
        self.assertEqual(payroll["50"], 120)

    def test_night_into_holiday(self):
        payroll = self.classify(
            date(2026, 11, 19),
            holidays={"2026-11-20"},
            check_in="19:30",
            check_out="20:00",
            break_start="21:00",
            break_end="07:00",
        )
        self.assertEqual(payroll["100"], 540)
        self.assertEqual(payroll["normal"], 90)

    def test_pairs_after_midnight_are_next_day(self):
        payroll = self.classify(
            date(2026, 10, 17),
            check_in="19:30",
            check_out="23:00",
            break_start="00:00",
            break_end="03:00",
            entry_3="04:00",
            exit_3="07:00",
        )
        self.assertEqual(payroll["50"], 150)
        self.assertEqual(payroll["100"], 420)

    def test_entry_rounded_to_next_midnight_uses_start_date(self):
        # Entrada sabado 23:50 arredondada para domingo 00:00: o app manda a
        # linha de sabado com start_date de domingo.
        payroll = self.classify(
            date(2026, 10, 17), check_in="00:00", check_out="07:00", start_date="2026-10-18"
        )
        self.assertEqual(payroll["100"], 420)
        self.assertEqual(payroll["50"], 0)
        # Sem start_date (ou fora do dia da linha e do seguinte), a data da linha.
        payroll = self.classify(date(2026, 10, 17), check_in="00:00", check_out="07:00")
        self.assertEqual(payroll["50"], 120)
        payroll = self.classify(
            date(2026, 10, 17), check_in="00:00", check_out="07:00", start_date="2026-10-25"
        )
        self.assertEqual(payroll["50"], 120)

    def test_html_intern_uses_period_holidays(self):
        record = main.ExportRecord(
            date="2026-11-19", check_in="19:30", check_out="20:00", break_start="21:00", break_end="07:00"
        )
        vals = main.classify_html_day(record, True, date(2026, 11, 19), False, None, {"2026-11-20"})
        self.assertEqual(vals["he100"], 540)


@needs_main
class OpenShiftStartTest(unittest.TestCase):
    def test_gestor_batch_night_is_closed(self):
        # Lote do gestor: saida 07:00 e entrada 19:30 na mesma data.
        records = [
            {"id": "out", "record_type": "check_out", "timestamp": "2026-10-15T07:00:00", "created_by_admin": True},
            {"id": "in", "record_type": "check_in", "timestamp": "2026-10-15T19:30:00", "created_by_admin": True},
        ]
        self.assertIsNone(main._open_shift_start(records))

    def test_app_check_in_after_gestor_exit_stays_open(self):
        records = [
            {"id": "out", "record_type": "check_out", "timestamp": "2026-10-15T07:00:00", "created_by_admin": True},
            {"id": "in", "record_type": "check_in", "timestamp": "2026-10-15T19:30:00"},
        ]
        self.assertEqual(main._open_shift_start(records)["id"], "in")

    def test_gestor_check_in_without_batch_exit_stays_open(self):
        records = [
            {"id": "out", "record_type": "check_out", "timestamp": "2026-10-15T07:00:00"},
            {"id": "in", "record_type": "check_in", "timestamp": "2026-10-15T19:30:00", "created_by_admin": True},
        ]
        self.assertEqual(main._open_shift_start(records)["id"], "in")

    def test_batch_night_on_day_with_app_shift_is_closed(self):
        records = [
            {"id": "nout", "record_type": "check_out", "timestamp": "2026-10-15T07:00:00", "created_by_admin": True},
            {"id": "in1", "record_type": "check_in", "timestamp": "2026-10-15T08:01:00"},
            {"id": "out1", "record_type": "check_out", "timestamp": "2026-10-15T17:02:00"},
            {"id": "nin", "record_type": "check_in", "timestamp": "2026-10-15T19:30:00", "created_by_admin": True},
        ]
        self.assertIsNone(main._open_shift_start(records))
        self.assertEqual(main._manual_night_exit_index(records, 3), 0)

    def test_batch_double_night_takes_exit_before_first_entry(self):
        manual = {"created_by_admin": True}
        records = [
            {"id": "nout", "record_type": "check_out", "timestamp": "2026-10-15T07:00:00", **manual},
            {"id": "in1", "record_type": "check_in", "timestamp": "2026-10-15T08:00:00", **manual},
            {"id": "out1", "record_type": "check_out", "timestamp": "2026-10-15T17:00:00", **manual},
            {"id": "nin", "record_type": "check_in", "timestamp": "2026-10-15T19:30:00", **manual},
        ]
        self.assertEqual(main._manual_night_exit_index(records, 3), 0)
        self.assertIsNone(main._manual_night_exit_index(records, 1))
        self.assertIsNone(main._open_shift_start(records))

    def test_batch_night_exit_after_day_shift_start(self):
        # 06:50-16:00 do app e a noite 19:00-07:00 do lote na mesma data: a
        # saida 07:00 cai depois da entrada 06:50 e a 16:00 a substitui.
        manual = {"created_by_admin": True}
        records = [
            {"id": "in1", "record_type": "check_in", "timestamp": "2026-10-14T06:50:00"},
            {"id": "nout", "record_type": "check_out", "timestamp": "2026-10-14T07:00:00", **manual},
            {"id": "out1", "record_type": "check_out", "timestamp": "2026-10-14T16:00:00"},
            {"id": "nin", "record_type": "check_in", "timestamp": "2026-10-14T19:00:00", **manual},
        ]
        self.assertEqual(main._manual_night_exit_index(records, 3), 1)
        self.assertIsNone(main._open_shift_start(records))
        # Tudo no lote: 06-15 e 19-07.
        batch = [
            {"id": "in1", "record_type": "check_in", "timestamp": "2026-10-14T06:00:00.000", **manual},
            {"id": "nout", "record_type": "check_out", "timestamp": "2026-10-14T07:00:00.000", **manual},
            {"id": "out1", "record_type": "check_out", "timestamp": "2026-10-14T15:00:00.000", **manual},
            {"id": "nin", "record_type": "check_in", "timestamp": "2026-10-14T19:00:00.000", **manual},
        ]
        self.assertEqual(main._manual_night_exit_index(batch, 3), 1)
        self.assertIsNone(main._open_shift_start(batch))
        # A saida do gestor que nenhuma outra substitui fecha o turno do dia.
        day_exit = [
            {"id": "in1", "record_type": "check_in", "timestamp": "2026-10-14T06:50:00"},
            {"id": "out1", "record_type": "check_out", "timestamp": "2026-10-14T16:00:00", **manual},
            {"id": "nin", "record_type": "check_in", "timestamp": "2026-10-14T19:00:00", **manual},
        ]
        self.assertIsNone(main._manual_night_exit_index(day_exit, 2))
        self.assertEqual(main._open_shift_start(day_exit)["id"], "nin")

    def test_batch_night_is_at_most_16_hours(self):
        # Entrada 08:00 do gestor no dia da saida 07:00 que corrige o
        # "Esqueci" da vespera: turno do dia aberto, nao noite de 23h.
        manual = {"created_by_admin": True}
        records = [
            {"id": "in15", "record_type": "check_in", "timestamp": "2026-10-15T19:30:00"},
            {"id": "forgot_in15", "record_type": "check_out", "timestamp": "2026-10-15T19:30:01"},
            {"id": "fix", "record_type": "check_out", "timestamp": "2026-10-16T07:00:00", **manual},
            {"id": "in16", "record_type": "check_in", "timestamp": "2026-10-16T08:00:00", **manual},
        ]
        self.assertIsNone(main._manual_night_exit_index(records, 3))
        self.assertEqual(main._open_shift_start(records)["id"], "in16")


@needs_main
class BatchRealDatesTest(unittest.TestCase):
    """Lote novo do gestor: data real e batch_real_dates. Vale como o app; as
    regras da noite do lote ficam so para o lote antigo."""

    legacy = {"created_by_admin": True}
    batch = {"created_by_admin": True, "batch_real_dates": True}

    def test_is_legacy_batch(self):
        self.assertTrue(main._is_legacy_batch({**self.legacy}))
        self.assertTrue(main._is_legacy_batch({**self.legacy, "batch_real_dates": False}))
        self.assertFalse(main._is_legacy_batch({**self.batch}))
        self.assertFalse(main._is_legacy_batch({}))
        self.assertFalse(main._is_legacy_batch({"batch_real_dates": True}))

    def test_new_batch_night_closes_on_next_day(self):
        records = [
            {"id": "in", "record_type": "check_in", "timestamp": "2026-10-15T19:30:00.000", **self.batch},
            {"id": "b1", "record_type": "break_start", "timestamp": "2026-10-16T02:00:00.000", **self.batch},
            {"id": "b2", "record_type": "break_end", "timestamp": "2026-10-16T03:00:00.000", **self.batch},
            {"id": "out", "record_type": "check_out", "timestamp": "2026-10-16T07:00:00.000", **self.batch},
        ]
        self.assertIsNone(main._manual_night_exit_index(records, 0))
        self.assertIsNone(main._open_shift_start(records))
        # So a entrada lancada: o turno fica aberto, como no app.
        self.assertEqual(main._open_shift_start(records[:1])["id"], "in")

    def test_flagged_pair_on_same_date_is_not_a_night(self):
        # Parece a noite do lote antigo, mas com a data real: saida solta e
        # entrada aberta.
        records = [
            {"id": "out", "record_type": "check_out", "timestamp": "2026-10-15T07:00:00.000", **self.batch},
            {"id": "in", "record_type": "check_in", "timestamp": "2026-10-15T19:30:00.000", **self.batch},
        ]
        self.assertIsNone(main._manual_night_exit_index(records, 1))
        self.assertEqual(main._open_shift_start(records)["id"], "in")
        # Mistura de lote antigo e novo tambem nao e noite.
        legacy_entry = [dict(records[0]), {**records[1], "batch_real_dates": None}]
        self.assertIsNone(main._manual_night_exit_index(legacy_entry, 1))
        self.assertEqual(main._open_shift_start(legacy_entry)["id"], "in")
        legacy_exit = [{**records[0], "batch_real_dates": None}, dict(records[1])]
        self.assertIsNone(main._manual_night_exit_index(legacy_exit, 1))
        self.assertEqual(main._open_shift_start(legacy_exit)["id"], "in")

    def test_new_batch_exit_fixes_forgotten_exit(self):
        records = [
            {"id": "in15", "record_type": "check_in", "timestamp": "2026-10-15T19:30:00"},
            {"id": "forgot_in15", "record_type": "check_out", "timestamp": "2026-10-15T19:30:01"},
            {"id": "fix", "record_type": "check_out", "timestamp": "2026-10-16T07:00:00.000", **self.batch},
        ]
        self.assertIsNone(main._open_shift_start(records))

    def test_new_batch_exit_is_not_taken_by_legacy_night(self):
        # Transicao: 15/10 pelo lote novo (saida em 16/10) e 16/10 pelo lote
        # antigo (saida da manha e entrada 19:30 em 16/10).
        for legacy_exit, night_exit in (("06:30", 1), ("07:30", 2)):
            records = sorted([
                {"id": "in15", "record_type": "check_in", "timestamp": "2026-10-15T19:30:00.000", **self.batch},
                {"id": "out15", "record_type": "check_out", "timestamp": "2026-10-16T07:00:00.000", **self.batch},
                {"id": "out16", "record_type": "check_out", "timestamp": f"2026-10-16T{legacy_exit}:00.000", **self.legacy},
                {"id": "in16", "record_type": "check_in", "timestamp": "2026-10-16T19:30:00.000", **self.legacy},
            ], key=lambda item: item["timestamp"])
            self.assertEqual(main._manual_night_exit_index(records, 3), night_exit)
            self.assertIsNone(main._open_shift_start(records))

    def test_day_shift_and_new_batch_night_on_same_date(self):
        records = [
            {"id": "in1", "record_type": "check_in", "timestamp": "2026-10-14T06:50:00"},
            {"id": "out1", "record_type": "check_out", "timestamp": "2026-10-14T16:00:00"},
            {"id": "nin", "record_type": "check_in", "timestamp": "2026-10-14T19:00:00.000", **self.batch},
        ]
        # Sem a saida lancada ainda, a noite fica aberta.
        self.assertEqual(main._open_shift_start(records)["id"], "nin")
        records.append(
            {"id": "nout", "record_type": "check_out", "timestamp": "2026-10-15T07:00:00.000", **self.batch}
        )
        self.assertIsNone(main._open_shift_start(records))


@needs_main
class StartRequestExpiredTest(unittest.TestCase):
    def test_request_just_before_midnight_is_still_valid(self):
        self.assertFalse(main._start_request_expired(datetime(2026, 10, 15, 23, 55), datetime(2026, 10, 16, 0, 3)))

    def test_old_request_expires(self):
        self.assertTrue(main._start_request_expired(datetime(2026, 10, 15, 8, 0), datetime(2026, 10, 15, 21, 0)))
        self.assertFalse(main._start_request_expired(None, datetime(2026, 10, 15, 21, 0)))

    def test_retry_window(self):
        first = datetime(2026, 10, 17, 8, 0)
        self.assertTrue(main._is_start_request_retry(first, first + timedelta(minutes=15)))
        self.assertFalse(main._is_start_request_retry(first, first + timedelta(hours=7)))
        self.assertTrue(main._is_start_request_retry(None, first))


class _FakeSnapshot:
    def __init__(self, store, collection, doc_id):
        self._store = store
        self._collection = collection
        self.id = doc_id
        self.reference = _FakeDocument(store, collection, doc_id)

    @property
    def exists(self):
        return self.id in self._store.setdefault(self._collection, {})

    def to_dict(self):
        data = self._store.setdefault(self._collection, {}).get(self.id)
        return dict(data) if data is not None else None


class _FakeDocument:
    def __init__(self, store, collection, doc_id):
        self._store = store
        self._collection = collection
        self.id = doc_id

    def get(self):
        return _FakeSnapshot(self._store, self._collection, self.id)

    def set(self, data):
        self._store.setdefault(self._collection, {})[self.id] = dict(data)

    def create(self, data):
        if self.id in self._store.setdefault(self._collection, {}):
            raise main.AlreadyExists("documento ja existe")
        self.set(data)

    def update(self, data):
        self._store[self._collection][self.id].update(data)

    def delete(self):
        self._store.setdefault(self._collection, {}).pop(self.id, None)


class _FakeQuery:
    _ids = itertools.count(1)

    def __init__(self, store, collection, filters=()):
        self._store = store
        self._collection = collection
        self._filters = tuple(filters)

    def where(self, field, op, value):
        return _FakeQuery(self._store, self._collection, self._filters + ((field, op, value),))

    @staticmethod
    def _matches(doc_id, data, field, op, value):
        # "__name__" compara o id do documento (o valor e uma referencia).
        if field == "__name__":
            current, value = doc_id, value.id
        else:
            current = data.get(field)
        if op == "==":
            return current == value
        if current is None:
            return False
        return current <= value if op == "<=" else current >= value

    def stream(self):
        for doc_id, data in list(self._store.setdefault(self._collection, {}).items()):
            if all(self._matches(doc_id, data, *item) for item in self._filters):
                yield _FakeSnapshot(self._store, self._collection, doc_id)

    def document(self, doc_id=None):
        return _FakeDocument(self._store, self._collection, doc_id or f"auto{next(self._ids)}")


class _FakeDb:
    def __init__(self):
        self.store = {}

    def collection(self, name):
        return _FakeQuery(self.store, name)


@needs_main
class SameTimestampOrderTest(unittest.TestCase):
    """Pontos no mesmo timestamp (o lote grava sem segundos) saem numa ordem
    fixa: fim da pausa, saida, entrada, inicio da pausa. Mesma ordem de
    compareRecords (frontend/lib/utils/shift_records.dart)."""

    batch = {"created_by_admin": True, "batch_real_dates": True}

    def test_sort_key_orders_types_at_same_timestamp(self):
        at = "2026-10-16T07:00:00.000"
        records = [
            {"id": "a", "record_type": "entry_3", "timestamp": at},
            {"id": "b", "record_type": "break_start", "timestamp": at},
            {"id": "c", "record_type": "check_in", "timestamp": at},
            {"id": "d", "record_type": "check_out", "timestamp": at},
            {"id": "e", "record_type": "exit_3", "timestamp": at},
            {"id": "f", "record_type": "break_end", "timestamp": at},
            {"id": "g", "record_type": "check_out", "timestamp": "2026-10-16T06:59:59.999"},
        ]
        expected = ["g", "e", "f", "d", "c", "a", "b"]
        for order in (records, list(reversed(records))):
            self.assertEqual([item["id"] for item in sorted(order, key=main._record_sort_key)], expected)

    def test_exit_and_next_entry_at_same_timestamp(self):
        # Noite 19:00-07:00 e dia a partir das 07:00 lancados no lote: a
        # saida da noite e a entrada do dia tem o mesmo timestamp. Com a
        # entrada primeiro, a saida fechava o dia e o turno aberto sumia.
        records = [
            {"id": "b", "user_id": "u1", "record_type": "check_in", "timestamp": "2026-10-16T07:00:00.000", **self.batch},
            {"id": "a", "user_id": "u1", "record_type": "check_out", "timestamp": "2026-10-16T07:00:00.000", **self.batch},
            {"id": "nin", "user_id": "u1", "record_type": "check_in", "timestamp": "2026-10-15T19:00:00.000", **self.batch},
        ]
        for order in (records, list(reversed(records))):
            db = _FakeDb()
            for record in order:
                db.store.setdefault("time_records", {})[record["id"]] = {
                    key: value for key, value in record.items() if key != "id"
                }
            ordered = main._user_records_sorted(db, "u1", lookback_days=0)
            self.assertEqual([item["id"] for item in ordered], ["nin", "a", "b"])
            self.assertEqual(main._open_shift_start(ordered)["id"], "b")

    def test_same_instant_with_and_without_fraction(self):
        # Saida do FECHAR PONTO gravada por este backend (isoformat sem a
        # fracao) e o fim da pausa do lote no mesmo horario: o tipo decide.
        exit_ = {"id": "out", "record_type": "check_out", "timestamp": "2026-10-15T19:30:00"}
        pause_end = {"id": "p2", "record_type": "exit_3", "timestamp": "2026-10-15T19:30:00.000", **self.batch}
        self.assertLess(main._record_sort_key(pause_end), main._record_sort_key(exit_))
        for order in ([exit_, pause_end], [pause_end, exit_]):
            self.assertEqual([item["id"] for item in main._sort_records(list(order))], ["p2", "out"])
        self.assertEqual(main._normalized_timestamp("2026-10-15T19:30"), "2026-10-15T19:30:00.000000")
        self.assertEqual(main._normalized_timestamp("2026-10-15T19:30:00.1234567"), "2026-10-15T19:30:00.123456")
        self.assertEqual(main._normalized_timestamp("invalido"), "invalido")


@needs_main
class BatchMinuteOrderTest(unittest.TestCase):
    """O ponto do lote (minuto cheio) no mesmo minuto de uma batida com
    segundos segue a ordem do tipo. Mesma regra de chronologicalOrder
    (frontend/lib/utils/shift_records.dart)."""

    batch = {"created_by_admin": True, "batch_real_dates": True}

    def punch(self, record_id, record_type, timestamp, user_id="u1", **extra):
        return {"id": record_id, "user_id": user_id, "record_type": record_type, "timestamp": timestamp, **extra}

    def lote(self, record_id, record_type, timestamp, user_id="u1"):
        return self.punch(record_id, record_type, timestamp, user_id, **self.batch)

    def assert_order(self, records, expected):
        for order in (records, list(reversed(records))):
            self.assertEqual([item["id"] for item in main._sort_records(list(order))], expected)

    def test_batch_entry_after_app_exit_in_same_minute(self):
        # Noite do app e dia do lote: a Entrada 07:00 do lote vinha antes da
        # Saida 07:00:14 e o turno do dia nao aparecia aberto.
        records = [
            self.punch("nin", "check_in", "2026-10-15T19:00:12.345000"),
            self.punch("nout", "check_out", "2026-10-16T07:00:14.512000"),
            self.lote("din", "check_in", "2026-10-16T07:00:00.000"),
        ]
        self.assert_order(records, ["nin", "nout", "din"])
        self.assertEqual(main._open_shift_start(main._sort_records(list(records)))["id"], "din")

    def test_batch_exit_before_app_entry_in_same_minute(self):
        records = [
            self.lote("nin", "check_in", "2026-10-15T19:00:00.000"),
            self.lote("nout", "check_out", "2026-10-16T07:00:00.000"),
            self.punch("din", "check_in", "2026-10-16T07:00:20.100000"),
        ]
        self.assert_order(records, ["nin", "nout", "din"])
        self.assertEqual(main._open_shift_start(main._sort_records(list(records)))["id"], "din")

    def test_batch_exit_after_app_pause_end_in_same_minute(self):
        records = [
            self.punch("in", "check_in", "2026-10-15T19:30:12"),
            self.punch("b1", "break_start", "2026-10-16T03:00:05"),
            self.punch("b2", "break_end", "2026-10-16T07:00:20"),
            self.lote("out", "check_out", "2026-10-16T07:00:00.000"),
            self.lote("p2", "exit_3", "2026-10-16T07:00:00.000"),
        ]
        self.assert_order(records, ["in", "b1", "p2", "b2", "out"])

    def test_other_user_in_same_minute_does_not_move_batch(self):
        records = [
            self.punch("other", "check_out", "2026-10-16T07:00:14", user_id="u2"),
            self.lote("din", "check_in", "2026-10-16T07:00:00.000"),
            self.punch("mine", "check_in", "2026-10-16T07:00:30"),
        ]
        self.assert_order(records, ["din", "other", "mine"])

    def test_forgotten_exit_stays_after_its_entry(self):
        records = [
            self.punch("in", "check_in", "2026-10-15T08:00:05"),
            self.punch("forgot_in", "check_out", "2026-10-15T08:00:06"),
            self.lote("p1", "break_start", "2026-10-15T08:00:00.000"),
        ]
        self.assert_order(records, ["in", "forgot_in", "p1"])

    def test_legacy_and_batch_points_follow_type_order(self):
        records = [
            self.lote("b", "check_in", "2026-10-16T07:00:00.000"),
            self.punch("l", "check_out", "2026-10-16T07:00:00.000", created_by_admin=True),
            self.lote("p", "break_end", "2026-10-16T07:00:00.000"),
        ]
        self.assert_order(records, ["p", "l", "b"])

    def test_offline_exit_ignores_batch_entry_later_in_same_minute(self):
        # Saida da noite batida offline as 07:00:14 e a Entrada 07:00 do lote
        # para o turno do dia: no momento da saida, o turno aberto e a noite.
        records = main._sort_records([
            self.punch("nin", "check_in", "2026-10-15T19:00:12.345000"),
            self.lote("din", "check_in", "2026-10-16T07:00:00.000"),
        ])
        before = main._records_before_punch(records, {
            "user_id": "u1", "record_type": "check_out", "timestamp": "2026-10-16T07:00:14.512000",
        })
        self.assertEqual([item["id"] for item in before], ["nin"])
        self.assertEqual(main._open_shift_start(before)["id"], "nin")
        later = main._records_before_punch(records, {
            "user_id": "u1", "record_type": "check_out", "timestamp": "2026-10-16T15:00:00",
        })
        self.assertEqual([item["id"] for item in later], ["nin", "din"])

    def test_offline_punch_keeps_record_of_same_instant_and_type(self):
        records = main._sort_records([
            self.punch("in", "check_in", "2026-10-15T08:00:00.000"),
            self.punch("out", "check_out", "2026-10-15T17:00:00"),
            self.punch("later", "check_in", "2026-10-15T19:00:00.000"),
        ])
        before = main._records_before_punch(records, {
            "user_id": "u1", "record_type": "check_out", "timestamp": "2026-10-15T17:00:00.000",
        })
        self.assertEqual([item["id"] for item in before], ["in", "out"])

    def test_offline_intern_exit_uses_night_shift_for_overtime(self):
        # Estagiario: a saida da noite batida offline as 07:00:14 chega
        # depois da Entrada 07:00 do lote (turno do dia). A hora extra e a do
        # turno da noite; pelo texto do horario, o turno aberto era o do dia
        # e a saida passava sem a liberacao.
        db = _FakeDb()
        user = {"id": "u1", "role": "intern", "name": "Estagiario", "email": "e@x"}
        patches = [
            mock.patch.object(main, "get_firestore_admin", return_value=db),
            mock.patch.object(main, "firestore", types.SimpleNamespace(SERVER_TIMESTAMP="ts")),
        ]
        for patcher in patches:
            patcher.start()
            self.addCleanup(patcher.stop)
        today = datetime.now(ZoneInfo("America/Sao_Paulo")).date()
        night = today - timedelta(days=2)
        morning = today - timedelta(days=1)
        records = [
            self.punch("nin", "check_in", f"{night.isoformat()}T19:00:12.345000"),
            self.lote("din", "check_in", f"{morning.isoformat()}T07:00:00.000"),
        ]
        for record in records:
            db.store.setdefault("time_records", {})[record["id"]] = {
                key: value for key, value in record.items() if key != "id"
            }
        payload = main.SchedulePunchRequest(
            record_type="check_out",
            timestamp=f"{morning.isoformat()}T07:00:14.512000",
            offline=True,
        )
        result = _call(main.schedule_controlled_punch_endpoint(payload, user))
        self.assertEqual(result["status"], "approval_required")
        self.assertEqual(result["request_type"], "overtime")


@needs_main
class StaleStartRequestTest(unittest.TestCase):
    """Pedido de entrada pendente de horas atras nao engole o novo, e liberar o
    antigo nao abre o turno na hora dele."""

    user = {"id": "u1", "role": "intern", "name": "Estagiario", "email": "e@x"}
    admin = {"id": "g1", "role": "manager", "name": "Gestor"}

    def setUp(self):
        self.db = _FakeDb()
        patches = [
            mock.patch.object(main, "get_firestore_admin", return_value=self.db),
            mock.patch.object(main, "firestore", types.SimpleNamespace(SERVER_TIMESTAMP="ts")),
            mock.patch.object(main, "_is_start_within_schedule", return_value=False),
            mock.patch.object(main, "_send_user_push", lambda *args, **kwargs: None),
        ]
        for patcher in patches:
            patcher.start()
            self.addCleanup(patcher.stop)
        self.now = datetime.now(ZoneInfo("America/Sao_Paulo")).replace(tzinfo=None, microsecond=0)
        self.day_key = self.now.date().isoformat()

    def add_request(self, doc_id, attempted, status="pending"):
        self.db.store.setdefault("work_schedule_approval_requests", {})[doc_id] = {
            "user_id": "u1",
            "request_type": "start_exception",
            "status": status,
            "attempted_at": attempted.isoformat(),
            "day_key": self.day_key,
            "created_at": attempted.isoformat(),
            "notes": "antes",
        }

    def ask(self):
        payload = main.ScheduleApprovalCreateRequest(
            request_type="start_exception", attempted_at=self.now.isoformat(), notes="de novo"
        )
        return _call(main.create_schedule_approval_request_endpoint(payload, self.user))

    def approve(self, request_id):
        payload = main.ScheduleApprovalDecisionRequest(reason="ok")
        return _call(main.approve_schedule_request_endpoint(request_id, payload, self.admin))

    def test_old_pending_request_is_superseded_by_new_attempt(self):
        self.add_request("old", self.now - timedelta(hours=7))
        response = self.ask()
        self.assertEqual(response["status"], "created")
        requests = self.db.store["work_schedule_approval_requests"]
        self.assertEqual(requests["old"]["status"], "superseded")
        # O gestor ve a hora do primeiro pedido no novo.
        self.assertEqual(
            response["request"]["previous_attempted_at"],
            (self.now - timedelta(hours=7)).isoformat(),
        )
        with self.assertRaises(main.HTTPException) as ctx:
            self.approve("old")
        self.assertEqual(ctx.exception.status_code, 409)
        self.assertNotIn("approval_old", self.db.store.get("time_records", {}))

    def test_retry_of_same_attempt_reuses_pending_request(self):
        self.add_request("recent", self.now - timedelta(minutes=5))
        response = self.ask()
        self.assertEqual(response["status"], "existing")
        self.assertEqual(response["request"]["id"], "recent")

    def test_approve_refuses_request_with_newer_attempt(self):
        # Dois pendentes (modo compativel nao marca o antigo).
        self.add_request("old", self.now - timedelta(hours=7))
        self.add_request("new", self.now)
        with self.assertRaises(main.HTTPException) as ctx:
            self.approve("old")
        self.assertEqual(ctx.exception.status_code, 409)
        self.assertEqual(self.approve("new")["status"], "approved")
        self.assertEqual(
            self.db.store["time_records"]["approval_new"]["timestamp"], self.now.isoformat()
        )


@needs_main
class FirestoreProfileReadTest(unittest.TestCase):
    """Falha do Firestore na leitura do perfil nao e falta de permissao."""

    def read_with_status(self, status):
        class Response:
            status_code = status

            def json(self):
                return {"fields": {}}

        class Client:
            def __init__(self, *args, **kwargs):
                pass

            async def __aenter__(self):
                return self

            async def __aexit__(self, *args):
                return False

            async def get(self, *args, **kwargs):
                return Response()

        with mock.patch.dict(os.environ, {"FIREBASE_PROJECT_ID": "p"}), \
                mock.patch.object(main, "_http_client", Client):
            return asyncio.run(main.get_firestore_document_with_user_token("users", "u1", "t"))

    def status_of(self, status):
        with self.assertRaises(main.HTTPException) as ctx:
            self.read_with_status(status)
        return ctx.exception.status_code

    def test_server_error_is_503(self):
        self.assertEqual(self.status_of(503), 503)
        self.assertEqual(self.status_of(500), 503)

    def test_permission_error_is_403(self):
        self.assertEqual(self.status_of(403), 403)
        self.assertEqual(self.status_of(401), 403)

    def test_ok_and_not_found(self):
        self.assertEqual(self.read_with_status(200), {})
        self.assertIsNone(self.read_with_status(404))


@needs_main
class ShiftDayTest(unittest.TestCase):
    today = date(2026, 10, 16)

    def test_accepts_recent_day(self):
        self.assertEqual(main._parse_shift_day("2026-10-15", self.today), date(2026, 10, 15))
        self.assertEqual(main._parse_shift_day("2026-10-02", self.today), date(2026, 10, 2))

    def test_rejects_invalid_future_or_old_day(self):
        for value in (None, "", "15/10/2026", "2026-10-17", "2026-10-01", "2026-13-01", "20261015"):
            self.assertIsNone(main._parse_shift_day(value, self.today), value)


@needs_main
class ReusableApprovalCheckInTest(unittest.TestCase):
    request = {"day_key": "2026-10-15", "attempted_at": "2026-10-15T21:00:00"}

    def records(self, *items):
        return [dict(item) for item in items]

    def test_reuses_check_in_of_same_request(self):
        records = self.records(
            {"id": "approval_r1", "record_type": "check_in", "timestamp": "2026-10-15T21:00:00",
             "schedule_approval_request_id": "r1"},
            {"id": "out", "record_type": "check_out", "timestamp": "2026-10-15T22:00:00"},
        )
        found = main._reusable_approval_check_in(records, "r1", self.request, None)
        self.assertEqual(found["id"], "approval_r1")

    def test_reuses_other_request_while_shift_is_open(self):
        records = self.records(
            {"id": "approval_r0", "record_type": "check_in", "timestamp": "2026-10-15T20:58:00",
             "schedule_approval_request_id": "r0"},
        )
        open_shift = main._open_shift_start(records)
        found = main._reusable_approval_check_in(records, "r1", self.request, open_shift)
        self.assertEqual(found["id"], "approval_r0")

    def test_reuses_other_request_when_asked_before_its_check_out(self):
        records = self.records(
            {"id": "approval_r0", "record_type": "check_in", "timestamp": "2026-10-15T20:58:00",
             "schedule_approval_request_id": "r0"},
            {"id": "out", "record_type": "check_out", "timestamp": "2026-10-15T23:00:00"},
        )
        found = main._reusable_approval_check_in(records, "r1", self.request, None)
        self.assertEqual(found["id"], "approval_r0")

    def test_new_entry_after_closed_shift_is_not_reused(self):
        records = self.records(
            {"id": "approval_r0", "record_type": "check_in", "timestamp": "2026-10-15T07:00:00",
             "schedule_approval_request_id": "r0"},
            {"id": "out", "record_type": "check_out", "timestamp": "2026-10-15T12:00:00"},
        )
        self.assertIsNone(main._reusable_approval_check_in(records, "r1", self.request, None))


@needs_main
class ExcelExportTest(unittest.TestCase):
    signature_texts = {
        "M8": "Assinaturas",
        "M21": "Assinatura do Funcionário",
        "M29": "Assinatura do Responsável",
        "M33": "Horas trabalhadas no mês",
    }

    def generate(self, **overrides):
        import openpyxl

        values = {
            "employee_name": "Teste",
            "employee_role": "employee",
            "month": "10",
            "year": "2026",
            "is_campo": True,
            "daily_records": [],
        }
        values.update(overrides)
        try:
            path = main.generate_excel_file(main.ExportData(**values))
        except FileNotFoundError as exc:
            self.skipTest(str(exc))
        self.addCleanup(os.unlink, path)
        workbook = openpyxl.load_workbook(path)
        self.addCleanup(workbook.close)
        return workbook.active

    def assert_signature_block(self, ws):
        for coord, text in self.signature_texts.items():
            self.assertEqual(ws[coord].value, text, coord)
        self.assertTrue(str(ws["M10"].value).startswith("Declaro que as informações"))
        self.assertEqual(ws["N34"].value, "=SUM(I9:I39)")
        self.assertEqual(ws["N34"].number_format, "[h]:mm")
        self.assertEqual(ws["I40"].value, "=SUM(I9:I39)")

    def test_campo_calendar_month(self):
        ws = self.generate(daily_records=[
            {"date": "2026-10-01", "check_in": "08:00", "check_out": "12:00"},
        ])
        self.assertEqual((ws["A9"].value, ws["B9"].value), ("01", "Qui"))
        self.assertEqual((ws["A11"].value, ws["B11"].value), ("03", "Sáb"))
        self.assertEqual((ws["A39"].value, ws["B39"].value), ("31", "Sáb"))
        self.assertEqual(ws["C9"].value, "08:00")
        self.assertEqual(ws["C6"].value, "Aux. Geral")
        self.assert_signature_block(ws)
        self.assertEqual(ws.print_area, "'Folha de Ponto'!$A$1:$P$39")
        self.assertLess(ws.column_dimensions["A"].width, main.EXCEL_CUSTOM_DAY_WIDTH)

    def test_intern_custom_period(self):
        ws = self.generate(
            employee_role="intern",
            is_campo=False,
            period_start="2026-09-29T00:00:00",
            period_end="2026-10-28T23:59:59",
            daily_records=[{"date": "2026-10-28", "check_in": "13:30", "check_out": "19:30"}],
        )
        self.assertEqual((ws["A9"].value, ws["B9"].value), ("29/09", "Ter"))
        # "29/09" nao cabe na largura do modelo.
        self.assertGreaterEqual(ws.column_dimensions["A"].width, main.EXCEL_CUSTOM_DAY_WIDTH)
        self.assertEqual((ws["A38"].value, ws["B38"].value), ("28/10", "Qua"))
        self.assertEqual(ws["C38"].value, "13:30")
        # 30 dias: a 31a linha fica vazia, sem formula.
        self.assertIsNone(ws["A39"].value)
        self.assertIsNone(ws["B39"].value)
        self.assertIsNone(ws["I39"].value)
        self.assertEqual(ws["B41"].value, "Normais")
        self.assert_signature_block(ws)
        self.assertEqual(ws.print_area, "'Folha de Ponto'!$A$1:$P$44")

    def test_exit_at_midnight_is_written_as_received(self):
        # MOD(D-C,1) ja conta a saida 00:00; trocar por 23:59 perdia 1 min.
        ws = self.generate(daily_records=[
            {"date": "2026-10-02", "check_in": "14:00", "check_out": "00:00"},
            {"date": "2026-10-03", "check_in": "19:30", "check_out": "00:00",
             "break_start": "01:00", "break_end": "07:00",
             "entry_3": "08:00", "exit_3": "00:00"},
        ])
        self.assertEqual(ws["D10"].value, "00:00")
        self.assertEqual(ws["D11"].value, "00:00")
        self.assertEqual(ws["H11"].value, "00:00")

    def test_month_without_dedicated_template(self):
        # Fevereiro de 2027 nao tinha modelo proprio: o modelo unico serve.
        ws = self.generate(month="02", year="2027")
        self.assertEqual((ws["A9"].value, ws["B9"].value), ("01", "Seg"))
        self.assertEqual((ws["A36"].value, ws["B36"].value), ("28", "Dom"))
        self.assertIsNone(ws["A37"].value)
        self.assertIsNone(ws["I37"].value)
        self.assert_signature_block(ws)


@needs_main
class DailyEntryTest(unittest.TestCase):
    """Modo diaria: uma diaria por pessoa por dia, pelo relogio do servidor."""

    user = {"id": "u1", "role": "employee", "name": "Diarista", "daily_mode": True}

    def setUp(self):
        self.db = _FakeDb()
        patches = [
            mock.patch.object(main, "get_firestore_admin", return_value=self.db),
            mock.patch.object(main, "firestore", types.SimpleNamespace(SERVER_TIMESTAMP="ts")),
        ]
        for patcher in patches:
            patcher.start()
            self.addCleanup(patcher.stop)
        self.now = main._local_now().replace(microsecond=0)

    def add(self, user=None, **payload):
        return _call(main.create_daily_entry_endpoint(
            main.DailyEntryRequest(**payload), user or self.user,
        ))

    def entries(self):
        return self.db.store.get("daily_entries", {})

    def test_only_daily_mode_users(self):
        for user in ({**self.user, "daily_mode": False}, {"id": "u2", "role": "employee"}):
            with self.assertRaises(main.HTTPException) as ctx:
                self.add(user)
            self.assertEqual(ctx.exception.status_code, 403)
        self.assertEqual(self.entries(), {})

    def test_one_entry_per_day(self):
        first = self.add()
        self.assertEqual(first["status"], "created")
        today = self.now.date().isoformat()
        self.assertEqual(first["date"], today)
        self.assertEqual(list(self.entries()), [f"u1_{today}"])
        # A hora mandada pelo app nao vale fora do offline: o dia e o do servidor.
        second = self.add(timestamp="2020-01-01T08:00:00")
        self.assertEqual(second["status"], "existing")
        self.assertEqual(second["registered_at"], first["registered_at"])
        self.assertEqual(len(self.entries()), 1)

    def test_offline_entry_counts_on_the_tap_day(self):
        yesterday = self.now - timedelta(days=1)
        response = self.add(offline=True, timestamp=yesterday.isoformat())
        self.assertEqual(response["status"], "created")
        self.assertEqual(response["date"], yesterday.date().isoformat())
        self.assertTrue(response["synced_from_offline"])
        # Reenvio da fila: a mesma diaria.
        self.assertEqual(self.add(offline=True, timestamp=yesterday.isoformat())["status"], "existing")

    def test_offline_entry_outside_window_is_refused(self):
        for at in (self.now - timedelta(days=8), self.now + timedelta(hours=1), None):
            with self.assertRaises(main.HTTPException) as ctx:
                self.add(offline=True, timestamp=at.isoformat() if at else None)
            self.assertEqual(ctx.exception.status_code, 400)
        self.assertEqual(self.entries(), {})

    def test_list_reads_only_the_period_of_the_user(self):
        store = self.db.store.setdefault("daily_entries", {})
        for user_id, day in (("u1", "2026-09-28"), ("u1", "2026-10-01"), ("u1", "2026-10-15"),
                             ("u1", "2026-10-29"), ("u2", "2026-10-02")):
            store[f"{user_id}_{day}"] = {"user_id": user_id, "date": day}

        def listing(user, **params):
            return _call(main.list_daily_entries_endpoint(user=user, **params))["entries"]

        own = listing(self.user, start="2026-09-29", end="2026-10-28")
        self.assertEqual([entry["date"] for entry in own], ["2026-10-01", "2026-10-15"])
        admin = {"id": "a1", "role": "admin"}
        other = listing(admin, start="2026-09-29", end="2026-10-28", user_id="u2")
        self.assertEqual([entry["id"] for entry in other], ["u2_2026-10-02"])
        for params in ({"user_id": "u2"}, {"start": "2026-10-28", "end": "2026-10-01"},
                       {"start": "2026-10-xx"}):
            with self.assertRaises(main.HTTPException):
                listing(self.user, **{"start": "2026-10-01", "end": "2026-10-28", **params})


@needs_main
class AdminDailyEntryTest(unittest.TestCase):
    """Admin lanca e apaga diarias, com motivo e auditoria."""

    admin = {"id": "a1", "role": "admin", "name": "Admin", "email": "a@x"}

    def setUp(self):
        self.db = _FakeDb()
        self.db.store["users"] = {"u1": {"name": "Diarista", "role": "employee", "daily_mode": True}}
        patches = [
            mock.patch.object(main, "get_firestore_admin", return_value=self.db),
            mock.patch.object(main, "firestore", types.SimpleNamespace(SERVER_TIMESTAMP="ts")),
        ]
        for patcher in patches:
            patcher.start()
            self.addCleanup(patcher.stop)
        self.today = main._local_now().date()

    def add(self, day, admin=None, user_id="u1", reason="esqueceu de marcar"):
        payload = main.AdminDailyEntryCreateRequest(user_id=user_id, date=day.isoformat(), reason=reason)
        return _call(main.admin_create_daily_entry_endpoint(payload, admin or self.admin))

    def delete(self, entry_id, admin=None, reason="marcou por engano"):
        payload = main.AdminDailyEntryDeleteRequest(reason=reason)
        return _call(main.admin_delete_daily_entry_endpoint(entry_id, payload, admin or self.admin))

    def status_of(self, call, *args, **kwargs):
        with self.assertRaises(main.HTTPException) as ctx:
            call(*args, **kwargs)
        return ctx.exception.status_code

    def audits(self):
        return list(self.db.store.get("system_audits", {}).values())

    def test_admin_adds_past_day_once(self):
        day = self.today - timedelta(days=3)
        response = self.add(day)
        entry_id = f"u1_{day.isoformat()}"
        self.assertEqual(response["id"], entry_id)
        stored = self.db.store["daily_entries"][entry_id]
        self.assertTrue(stored["created_by_admin"])
        self.assertEqual((stored["user_name"], stored["admin_reason"]), ("Diarista", "esqueceu de marcar"))
        self.assertEqual(self.audits()[0]["action"], "admin_add_daily_entry")
        # O mesmo dia de novo: ja existe.
        self.assertEqual(self.status_of(self.add, day), 409)

    def test_add_refuses_future_unknown_user_blank_reason_and_manager(self):
        self.assertEqual(self.status_of(self.add, self.today + timedelta(days=1)), 400)
        self.assertEqual(self.status_of(self.add, self.today, user_id="nao-existe"), 404)
        self.assertEqual(self.status_of(self.add, self.today, reason="  "), 400)
        manager = {"id": "g1", "role": "manager", "name": "Gestor"}
        self.assertEqual(self.status_of(self.add, self.today, admin=manager), 403)
        self.assertNotIn("daily_entries", {k for k, v in self.db.store.items() if v})

    def test_admin_deletes_with_audit(self):
        self.add(self.today)
        entry_id = f"u1_{self.today.isoformat()}"
        self.assertEqual(self.delete(entry_id)["status"], "deleted")
        self.assertNotIn(entry_id, self.db.store["daily_entries"])
        audit = self.audits()[-1]
        self.assertEqual((audit["action"], audit["reason"]), ("admin_delete_daily_entry", "marcou por engano"))
        self.assertEqual(audit["entry"]["id"], entry_id)
        self.assertEqual(self.status_of(self.delete, entry_id), 404)

    def test_delete_needs_admin_and_reason(self):
        self.add(self.today)
        entry_id = f"u1_{self.today.isoformat()}"
        manager = {"id": "g1", "role": "manager", "name": "Gestor"}
        self.assertEqual(self.status_of(self.delete, entry_id, admin=manager), 403)
        self.assertEqual(self.status_of(self.delete, entry_id, reason=""), 400)
        self.assertIn(entry_id, self.db.store["daily_entries"])


@needs_main
class DailyStandardDayTest(unittest.TestCase):
    """Dia padrao da diaria. Os mesmos casos estao em
    frontend/test/daily_entries_test.dart: a Dash e a folha dao o mesmo horario."""

    cases = [
        ("Diarista Teste", date(2026, 10, 1), 8, [("08:00", "12:00"), ("13:00", "17:00")]),
        ("José Ação", date(2026, 10, 17), 8, [("08:00", "12:00"), ("13:15", "17:15")]),
        ("José Ação", date(2026, 10, 18), 7.5, [("08:30", "12:30"), ("13:45", "17:15")]),
        ("Ana", date(2026, 10, 19), 4, [("08:30", "12:30")]),
    ]

    def test_reference_days(self):
        for name, day, hours, expected in self.cases:
            sessions = main.daily_standard_sessions(name, day, hours)
            self.assertEqual(
                [(start.strftime("%H:%M"), end.strftime("%H:%M")) for start, end in sessions],
                expected, (name, day),
            )


@needs_main
class DailyModeExportTest(unittest.TestCase):
    """Exportacao do modo diaria: aba com os dias e a folha padrao."""

    def data(self, **overrides):
        values = {
            "employee_name": "Diarista Teste",
            "employee_role": "employee",
            "month": "10",
            "year": "2026",
            "is_campo": True,
            "daily_records": [],
            "daily_mode": True,
            # 30/09 fica fora do mes; a data repetida conta uma vez.
            "daily_dates": ["2026-10-05", "2026-10-01", "2026-10-05", "2026-09-30", "2026-10-18"],
            "daily_hours": 8.0,
        }
        values.update(overrides)
        return main.ExportData(**values)

    def workbook(self, **overrides):
        import openpyxl

        try:
            path = main.generate_excel_file(self.data(**overrides))
        except FileNotFoundError as exc:
            self.skipTest(str(exc))
        self.addCleanup(os.unlink, path)
        workbook = openpyxl.load_workbook(path)
        self.addCleanup(workbook.close)
        return workbook

    def test_standard_day_on_each_daily_entry(self):
        records = main.daily_mode_records(self.data())
        self.assertEqual([record.date for record in records], ["2026-10-01", "2026-10-05", "2026-10-18"])
        for record in records:
            self.assertIn(record.check_in, ("08:00", "08:30"))
            pair = main.ExportRecord(date=record.date, check_in=record.check_in, check_out=record.check_out,
                                     break_start=record.break_start, break_end=record.break_end)
            self.assertEqual(main.session_minutes(pair), 480)
        # Mesmo nome e mesma data: mesmos horarios em toda exportacao.
        again = main.daily_mode_records(self.data())
        self.assertEqual([r.model_dump() for r in records], [r.model_dump() for r in again])

    def test_short_or_invalid_daily_hours(self):
        short = main.daily_mode_records(self.data(daily_hours=4))
        self.assertTrue(all(r.break_start is None and main.session_minutes(r) == 240 for r in short))
        for hours in (None, 0, 30):
            default = main.daily_mode_records(self.data(daily_hours=hours))
            self.assertTrue(all(main.session_minutes(r) == 480 for r in default))

    def test_excel_has_daily_sheet_and_timesheet(self):
        workbook = self.workbook()
        self.assertEqual(workbook.sheetnames, ["Diárias", "Folha de Ponto"])
        daily = workbook["Diárias"]
        timesheet = workbook["Folha de Ponto"]
        self.assertEqual(workbook.active.title, "Diárias")
        self.assertTrue(daily.sheet_view.tabSelected)
        self.assertFalse(timesheet.sheet_view.tabSelected)

        self.assertEqual(daily["A3"].value, "Colaborador: Diarista Teste")
        self.assertEqual(daily["A4"].value, "Função: Aux. Geral")
        self.assertEqual(daily["A5"].value, "Período: Outubro/2026")
        rows = [
            (daily.cell(row=r, column=1).value, daily.cell(row=r, column=2).value.date(),
             daily.cell(row=r, column=3).value)
            for r in range(8, 11)
        ]
        self.assertEqual(rows, [
            (1, date(2026, 10, 1), "Quinta-feira"),
            (2, date(2026, 10, 5), "Segunda-feira"),
            (3, date(2026, 10, 18), "Domingo"),
        ])
        self.assertEqual((daily["A12"].value, daily["C12"].value), ("Total de diárias", 3))
        self.assertEqual(daily.print_area, "'Diárias'!$A$1:$C$12")

        # Folha padrao: horario nos dias com diaria (linha 9 = dia 01) e so neles.
        self.assertIn(timesheet["C9"].value, ("08:00", "08:30"))
        self.assertIsNotNone(timesheet["F9"].value)
        self.assertIn(timesheet["C13"].value, ("08:00", "08:30"))
        self.assertIn(timesheet["C26"].value, ("08:00", "08:30"))
        self.assertIsNone(timesheet["C10"].value)
        self.assertEqual(timesheet["N34"].value, "=SUM(I9:I39)")
        self.assertEqual(timesheet.print_area, "'Folha de Ponto'!$A$1:$P$39")

    def test_excel_without_entries(self):
        workbook = self.workbook(daily_dates=[])
        daily = workbook["Diárias"]
        self.assertEqual((daily["A9"].value, daily["C9"].value), ("Total de diárias", 0))
        self.assertIsNone(workbook["Folha de Ponto"]["C9"].value)

    def test_filename_and_html(self):
        data = self.data()
        self.assertEqual(main.excel_download_filename(data), "Folha_Ponto_Diarista Teste_10_diarias.xlsx")
        path = main.generate_batch_html_file([data])
        self.addCleanup(os.unlink, path)
        with open(path, encoding="utf-8") as handle:
            content = handle.read()
        self.assertIn('<span class="badge diaria">Diária</span>', content)
        self.assertIn("Modo diaria - 3 diaria(s) no periodo", content)


@needs_main
class AdminDailyModeUpdateTest(unittest.TestCase):
    admin = {"id": "a1", "role": "admin", "name": "Admin", "email": "a@x"}

    def setUp(self):
        self.db = _FakeDb()
        self.db.store["users"] = {"u1": {"name": "Pessoa", "role": "employee", "daily_mode": True}}

        class AuthError(Exception):
            pass

        auth = types.SimpleNamespace(
            update_user=lambda *args, **kwargs: None,
            EmailAlreadyExistsError=AuthError,
            UserNotFoundError=AuthError,
        )
        patches = [
            mock.patch.object(main, "get_firestore_admin", return_value=self.db),
            mock.patch.object(main, "firestore", types.SimpleNamespace(SERVER_TIMESTAMP="ts")),
            mock.patch.object(main, "firebase_auth", auth),
        ]
        for patcher in patches:
            patcher.start()
            self.addCleanup(patcher.stop)

    def save(self, role="employee", **extra):
        payload = main.AdminUserUpdateRequest(
            name="Pessoa", email="p@x.com", role=role, job_title="Aux. Geral",
            daily_hours=8.0, cpf="", **extra,
        )
        _call(main.update_user_admin_endpoint("u1", payload, self.admin))
        return self.db.store["users"]["u1"]

    def test_daily_mode_is_saved(self):
        self.assertFalse(self.save(daily_mode=False)["daily_mode"])
        self.assertTrue(self.save(daily_mode=True)["daily_mode"])

    def test_old_app_without_field_keeps_it(self):
        self.assertTrue(self.save()["daily_mode"])

    def test_intern_never_in_daily_mode(self):
        self.assertFalse(self.save(role="intern", daily_mode=True)["daily_mode"])


# CPF valido de teste (digitos verificadores certos).
VALID_CPF = "52998224725"


class _FakeAuth:
    """firebase_auth de mentira: contas por uid, com e-mail e senha."""

    class EmailAlreadyExistsError(Exception):
        pass

    class UserNotFoundError(Exception):
        pass

    def __init__(self):
        self.accounts = {}
        self._ids = itertools.count(1)

    def create_user(self, email, password, display_name=None):
        if any(account["email"] == email for account in self.accounts.values()):
            raise self.EmailAlreadyExistsError(email)
        uid = f"new{next(self._ids)}"
        self.accounts[uid] = {"email": email, "password": password, "disabled": False}
        return types.SimpleNamespace(uid=uid)

    def get_user(self, uid):
        if uid not in self.accounts:
            raise self.UserNotFoundError(uid)
        return types.SimpleNamespace(uid=uid, email=self.accounts[uid]["email"])

    def update_user(self, uid, **fields):
        if uid not in self.accounts:
            raise self.UserNotFoundError(uid)
        self.accounts[uid].update(fields)

    def delete_user(self, uid):
        self.accounts.pop(uid, None)


@needs_main
class CpfRulesTest(unittest.TestCase):
    def test_valid_and_invalid_cpf(self):
        self.assertTrue(main.is_valid_cpf(VALID_CPF))
        for value in ("52998224724", "11111111111", "5299822472", "", "abcdefghijk"):
            self.assertFalse(main.is_valid_cpf(value), value)
        self.assertEqual(main.cpf_digits("529.982.247-25"), VALID_CPF)
        self.assertEqual(main.format_cpf(VALID_CPF), "529.982.247-25")

    def test_login_emails(self):
        self.assertEqual(main._login_emails(" Ana@X.com ", VALID_CPF), ("ana@x.com", "ana@x.com"))
        self.assertEqual(
            main._login_emails("", VALID_CPF),
            ("", f"{VALID_CPF}@{main.CPF_LOGIN_EMAIL_DOMAIN}"),
        )
        for email, digits in (("sem-arroba", VALID_CPF), ("", "123")):
            with self.assertRaises(main.HTTPException):
                main._login_emails(email, digits)


@needs_main
class CpfAccessTest(unittest.TestCase):
    """Login por CPF, cadastro pelo admin e troca de senha no primeiro acesso."""

    admin = {"id": "a1", "role": "admin", "name": "Admin", "email": "a@x"}

    def setUp(self):
        self.db = _FakeDb()
        self.auth = _FakeAuth()
        self.passwords_ok = {}
        main._login_failures.clear()

        async def check_password(email, password):
            account = next(a for a in self.auth.accounts.values() if a["email"] == email)
            if account.get("disabled"):
                return "disabled"
            return "ok" if account["password"] == password else "invalid"

        patches = [
            mock.patch.object(main, "get_firestore_admin", return_value=self.db),
            mock.patch.object(main, "firestore", types.SimpleNamespace(SERVER_TIMESTAMP="ts")),
            mock.patch.object(main, "firebase_auth", self.auth),
            mock.patch.object(main, "_check_firebase_password", check_password),
        ]
        for patcher in patches:
            patcher.start()
            self.addCleanup(patcher.stop)
        self.addCleanup(main._login_failures.clear)

    def create(self, admin=None, **fields):
        payload = main.AdminUserCreateRequest(**{"name": "Pessoa Nova", "cpf": "529.982.247-25", **fields})
        return _call(main.create_user_admin_endpoint(payload, admin or self.admin))

    def login(self, cpf, password, ip="1.2.3.4"):
        request = types.SimpleNamespace(headers={"x-real-ip": ip}, client=None)
        return _call(main.cpf_login_endpoint(main.CpfLoginRequest(cpf=cpf, password=password), request))

    def status_of(self, call, *args, **kwargs):
        with self.assertRaises(main.HTTPException) as ctx:
            call(*args, **kwargs)
        return ctx.exception.status_code, ctx.exception.detail

    def test_admin_creates_user_with_default_password(self):
        response = self.create(daily_mode=True)
        uid = response["id"]
        self.assertEqual((response["cpf"], response["default_password"]), ("529.982.247-25", "123456"))
        profile = self.db.store["users"][uid]
        self.assertTrue(profile["must_change_password"])
        self.assertTrue(profile["daily_mode"])
        self.assertEqual((profile["email"], profile["cpf_digits"], profile["job_title"]), ("", VALID_CPF, "Aux. Geral"))
        # Sem e-mail, a conta usa o interno do CPF.
        self.assertEqual(self.auth.accounts[uid]["email"], f"{VALID_CPF}@{main.CPF_LOGIN_EMAIL_DOMAIN}")
        self.assertEqual(self.auth.accounts[uid]["password"], "123456")

    def test_create_refuses_duplicate_invalid_cpf_and_manager(self):
        self.create()
        self.assertEqual(self.status_of(self.create, cpf="52998224725")[0], 409)
        self.assertEqual(self.status_of(self.create, cpf="12345678900")[0], 400)
        manager = {"id": "g1", "role": "manager", "name": "Gestor"}
        self.assertEqual(self.status_of(self.create, admin=manager, cpf="11144477735")[0], 403)
        self.assertEqual(len(self.auth.accounts), 1)

    def test_login_by_cpf_in_any_format(self):
        uid = self.create(email="pessoa@x.com")["id"]
        self.assertEqual(self.login("52998224725", "123456"), {"email": "pessoa@x.com"})
        # Cadastro antigo gravava "529.982.247.25".
        self.db.store["users"][uid]["cpf"] = "529.982.247.25"
        self.assertEqual(self.login("529.982.247-25", "123456"), {"email": "pessoa@x.com"})

    def test_wrong_password_and_unknown_cpf_look_the_same(self):
        self.create()
        wrong = self.status_of(self.login, VALID_CPF, "errada")
        unknown = self.status_of(self.login, "11144477735", "123456")
        self.assertEqual(wrong, unknown)
        self.assertEqual(wrong[0], 401)

    def test_disabled_and_duplicate(self):
        uid = self.create()["id"]
        self.auth.accounts[uid]["disabled"] = True
        self.assertEqual(self.status_of(self.login, VALID_CPF, "123456")[0], 403)
        self.db.store["users"]["outro"] = {"cpf": VALID_CPF}
        self.assertEqual(self.status_of(self.login, VALID_CPF, "123456")[0], 409)

    def test_too_many_failures_block_for_a_while(self):
        self.create()
        for _ in range(main._LOGIN_MAX_FAILURES):
            self.status_of(self.login, VALID_CPF, "errada")
        # Bloqueado mesmo com a senha certa, e de outro aparelho (por CPF).
        self.assertEqual(self.status_of(self.login, VALID_CPF, "123456", ip="9.9.9.9")[0], 429)

    def change(self, user, password):
        payload = main.OwnPasswordChangeRequest(new_password=password)
        return _call(main.change_own_password_endpoint(payload, user))

    def test_first_access_password_change(self):
        uid = self.create()["id"]
        user = {"id": uid, **self.db.store["users"][uid]}
        for bad in ("123456", "12345", VALID_CPF):
            self.assertEqual(self.status_of(self.change, user, bad)[0], 400)
        self.assertEqual(self.change(user, "minhaSenha9")["status"], "success")
        self.assertEqual(self.auth.accounts[uid]["password"], "minhaSenha9")
        self.assertFalse(self.db.store["users"][uid]["must_change_password"])
        self.assertEqual(self.login(VALID_CPF, "minhaSenha9"), {"email": f"{VALID_CPF}@{main.CPF_LOGIN_EMAIL_DOMAIN}"})

    def update(self, uid, **fields):
        values = {
            "name": "Pessoa Nova", "email": "", "role": "employee", "job_title": "Aux. Geral",
            "daily_hours": 8.0, "cpf": "529.982.247-25",
        }
        values.update(fields)
        _call(main.update_user_admin_endpoint(uid, main.AdminUserUpdateRequest(**values), self.admin))
        return self.db.store["users"][uid]

    def test_admin_reset_password_forces_change(self):
        uid = self.create()["id"]
        self.change({"id": uid, **self.db.store["users"][uid]}, "minhaSenha9")
        profile = self.update(uid, new_password="123456")
        self.assertTrue(profile["must_change_password"])
        self.assertEqual(self.auth.accounts[uid]["password"], "123456")

    def test_admin_edit_email_and_cpf(self):
        uid = self.create()["id"]
        # E-mail de verdade passa a ser o login da conta; tirando, volta o interno.
        self.update(uid, email="Pessoa@X.com")
        self.assertEqual(self.auth.accounts[uid]["email"], "pessoa@x.com")
        profile = self.update(uid, email="", cpf="111.444.777-35")
        self.assertEqual(self.auth.accounts[uid]["email"], f"11144477735@{main.CPF_LOGIN_EMAIL_DOMAIN}")
        self.assertEqual((profile["cpf"], profile["cpf_digits"], profile["email"]), ("111.444.777-35", "11144477735", ""))
        # Editar sem senha nova nao mexe na troca pendente do primeiro acesso.
        self.assertTrue(profile["must_change_password"])
        # CPF de outro cadastro, ou invalido, nao.
        self.db.store["users"]["outro"] = {"cpf": VALID_CPF}
        self.assertEqual(self.status_of(self.update, uid, cpf=VALID_CPF)[0], 409)
        self.assertEqual(self.status_of(self.update, uid, cpf="12345678900")[0], 400)


if __name__ == "__main__":
    unittest.main()
